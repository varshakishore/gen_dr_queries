#!/usr/bin/env python3
"""
Regenerate the annotated hard questions with the current make-harder prompt.

Rebuilds the exact row set of the "random assignment" sheet in
`Analyse gen hard questions.xlsx`, re-runs ONLY the make-harder step for each row,
and writes a new workbook with the same shape so the two can be annotated the same way.

WHAT IS HELD FIXED
------------------
Each annotated question is traced back to the `claude_call` log record that produced it
(`purpose="harder"`, matched on the generated `updated_question`). That record carries the
verbatim system prompt and user message of the original call, so the regeneration reuses:

  * the same system prompt -- same answering-system profile, same strategy menu, same
    banned-strategy list, same round-specific few-shot examples;
  * the same user message -- including the prior-attempt feedback block for round-1 rows;
  * the same strategy, pinned explicitly (see --no-pin-strategy to let the model re-choose).

The ONLY difference is the change under test: the `required_reasoning_process` key is
inserted into the OUTPUT FORMAT block and its RULE line into the RULES block, exactly as
`research_pipeline.py` now does. Few-shot examples inside the recovered prompt are left
alone -- the originals predate the field, and that mirrors what a real round-N run sees.

The answering system is NOT run: no research-server call, no judge, no criterion
verification. Only the generator step.

ROW SET
-------
Reproduces the sheet's own selection rule, verified to match all 30 rows and all 6
annotator assignments: take `verified.kept.json` in file order, keep entries with
`criterion_rewritten == False`, take the first `--n-rows`, and deal them round-robin to
the annotators in `--annotators` order. Rows are then grouped by annotator, as in the
source sheet.

Generations are checkpointed to `--generations-json` as they complete; re-running reuses
that file and skips the API, so building the workbook is free after the first pass.

Examples:
  python regenerate_annotation_sheet.py \\
      --xlsx "~/Downloads/Analyse gen hard questions.xlsx" \\
      --kept runs/loop_250/verified.kept.json \\
      --logs runs/loop_250 \\
      --generations-json regen/generations.json \\
      --out regen/regenerated.xlsx

Requires ANTHROPIC_API_KEY.
"""

import argparse
import json
import os
import re
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import llm_client
import research_pipeline as RP

DEFAULT_ANNOTATORS = ["Varsha", "Jena", "Lucy", "Jay", "Bingbing", "Doug"]

# The insertions the prompt change makes. Anchors are the lines they follow; each anchor
# occurs exactly once in a rendered make-harder prompt (the few-shot blocks carry real
# values, not these angle-bracket placeholders). Keep the inserts byte-identical to
# research_pipeline's templates -- the assert below fails the run if they drift.
FORMAT_ANCHOR = '  "updated_question": "<the rewritten question>",'
FORMAT_INSERT = ('  "required_reasoning_process": ["<a reasoning operation the system '
                 'must perform>", "<another>", "..."],')
RULES_ANCHOR = ("- The updated question length should change by fewer than 15 words "
                "from the seed.")
RULES_INSERT = (
    "- First work out required_reasoning_process: the reasoning OPERATIONS a system must perform to get this question right — the kind of retrieval, comparison, adjudication or inference the question demands, not the conclusions it should reach. Keep it SPARSE: typically 2-4 entries, each naming a distinct kind of reasoning work, in the order it has to happen. Write \"reconcile findings that conflict across study populations\", not \"state that study A and study B disagree\". If an entry could be satisfied by copying a sentence out of a single retrieved source, it is a content requirement, not a reasoning operation — move it to the verification criterion.\n"
    "- Name the KIND OF SEARCH and the KIND OF ANALYSIS each operation takes. Searches differ: a targeted lookup, a sweep across subfields, a hunt for disconfirming evidence, a search for a controlled comparison specifically, a search whose informative outcome is that nothing exists, a search of a neighbouring literature that never uses the question's vocabulary. So do analyses: reconciling conflicting results, causal identification, construct-validity checking, base-rate or denominator reasoning, adjudication between incompatible frameworks, composing facts no single source combines. Say which, rather than writing a generic \"retrieve relevant evidence\". Only spell this out where it genuinely bears on getting the question right: if an ordinary lookup is all a step takes, say that plainly. Do not invent an exotic search or analysis the question does not call for, and do not add an operation just to name a search or analysis type — a question needing one operation gets one entry.\n"
    "- Consider whether some NON-OBVIOUS search would fully answer the question — an unusual query formulation, an adjacent field that studies the same thing under another name, or a review, registry or dataset that already did the work. Many questions have no such shortcut; do not manufacture one. But if one exists, the question is retrievable rather than hard: either revise the question so that no single search resolves it, or name that search in required_reasoning_process and make the criterion turn on what it still leaves unresolved.\n"
    "- Then DERIVE the verification criterion. It MUST enforce the reasoning process: an answer that skipped, faked, or botched the required reasoning has to FAIL even when it states a plausible-sounding conclusion. Beyond that, the criterion is where correctness lives — it must also pin down the specific entities, claims, distinctions and failure modes that make an answer right or wrong for THIS question. The reasoning process says HOW the answer must be reached; the criterion says WHAT must be true of it.\n"
    "- CONCISENESS IS A HARD REQUIREMENT of the criterion, and it governs every rule below. Every word must change whether some answer passes or fails: if deleting a phrase would leave the same answers passing and the same answers failing, delete it. Name a class rather than listing its members, state a check once rather than restating it, and never elaborate or illustrate to make a check feel thorough. A shorter criterion that decides the same verdicts is always the better criterion.\n"
    "- State as many distinct checks as the question actually needs — there is no fixed number and no one-check-per-step correspondence; a single step may need several checks, and several steps may collapse into one. You may add additional specific requirements, but they must clearly and directly arise from the question or the reasoning process.\n"
    "- Write the criterion as a bulleted checklist, not prose: the line \"Answer must establish:\", then one \"- \" bullet per check. Bullets are fragments rather than sentences. A bullet runs as long as it must to pin down its one concept and no longer; length comes from the number of checks, never from elaborating one.\n"
    "- One bullet per check. Keep a compound in a SINGLE bullet when the relation between its parts is what is being checked — \"both X and Y\", \"connect X to Y\", \"X does not follow without Y\" — since splitting those lets an answer satisfy one half and pass. Split only requirements that are independently checkable.\n"
    "- Do not write bullets about the STATUS of the answer's conclusion — no \"conditional conclusion: ...\", no \"calibration that ...\", no \"appropriate qualification of ...\". A qualification belongs inside the bullet stating the claim it qualifies, not in a bullet of its own.\n"
    "- Add a \"Fails if ...\" bullet only for a wrong answer no bullet above already excludes — typically the most tempting one. Write several if several are needed, or none at all when the checks are already sufficient. Never restate a bullet as a \"Fails if\"."
)

assert (FORMAT_INSERT in RP.PROMPT_TO_MAKE_HARDER_QUESTION_EXPLORE
        and RULES_INSERT in RP.PROMPT_TO_MAKE_HARDER_QUESTION_EXPLORE), \
    "FORMAT_INSERT/RULES_INSERT no longer match research_pipeline's make-harder prompt"

STRATEGY_PIN = (
    "\n\nSTRATEGY CONSTRAINT: Use exactly the strategy below to make the question harder. "
    "Do not select or invent a different one; this run is a controlled comparison against "
    "an earlier question built with this same strategy.\n{strategy}\n"
)


class LockedRunLogger(RP.RunLogger):
    """RunLogger guarded by a lock; the generation calls run concurrently."""

    def __init__(self, *a, **kw):
        self._lock = threading.Lock()   # before super(): RunLogger.__init__ already logs
        super().__init__(*a, **kw)

    def _write(self, record: dict) -> None:
        with self._lock:
            super()._write(record)


def norm(s: str) -> str:
    """Loose key for matching question text across the sheet, the logs and the JSON."""
    return re.sub(r"\W+", " ", (s or "").lower()).strip()


def col_index(header: list, label: str) -> int:
    """Index of the first header cell containing `label`."""
    idx = next((i for i, h in enumerate(header) if h and label in str(h)), None)
    assert idx is not None, f"no column matching {label!r} in header: {header}"
    return idx


# ---------------------------------------------------------------------------
# Row set
# ---------------------------------------------------------------------------

def build_assignment(kept_path: Path, annotators: list, n_rows: int) -> list:
    """The sheet's row set: eligible entries in file order, dealt round-robin."""
    kept = json.loads(kept_path.read_text())
    eligible = [e for e in kept if not e.get("criterion_rewritten")]
    assert len(eligible) >= n_rows, (
        f"{kept_path}: only {len(eligible)} entries with criterion_rewritten==False, "
        f"need {n_rows}"
    )
    rows = []
    for pos, entry in enumerate(eligible[:n_rows]):
        rows.append({
            "position": pos,
            "annotator": annotators[pos % len(annotators)],
            "seed_question": entry["seed_question"],
            "original_question": entry["question"],
            "original_strategy": entry["strategy"],
            "original_criterion": entry["criterion"],
            "source_run": entry.get("source_run", ""),
            "round": entry.get("round"),
        })
    return rows


def verify_against_sheet(rows: list, xlsx: Path, sheet: str) -> None:
    """Assert the reproduced row set matches the sheet it is meant to mirror."""
    import openpyxl

    ws = openpyxl.load_workbook(xlsx, read_only=True)[sheet]
    grid = list(ws.iter_rows(values_only=True))
    header = grid[0]
    a_col = col_index(header, "Annotator")
    q_col = col_index(header, "Final Question")
    sheet_pairs = [(str(r[a_col]).strip(), norm(r[q_col]))
                   for r in grid[1:] if r[q_col] and str(r[q_col]).strip()]
    mine = sorted((r["annotator"], norm(r["original_question"])) for r in rows)
    assert sorted(sheet_pairs) == mine, (
        "reproduced row set does not match the sheet; "
        f"{len(sheet_pairs)} sheet rows vs {len(rows)} reproduced"
    )
    print(f"[rows] reproduced all {len(rows)} rows and annotator assignments from '{sheet}'")


# ---------------------------------------------------------------------------
# Prompt recovery
# ---------------------------------------------------------------------------

def index_harder_calls(log_root: Path) -> dict:
    """Map normalized generated question -> the claude_call record that produced it."""
    out: dict = {}
    files = sorted(log_root.rglob("run-sample_*.jsonl"))
    assert files, f"{log_root}: no run-sample_*.jsonl logs found"
    for f in files:
        for line in f.open():
            rec = json.loads(line)
            if (rec.get("kind") != "claude_call" or rec.get("purpose") != "harder"
                    or rec.get("error")):
                continue
            try:
                data = RP.extract_json(rec["response_text"])
            except ValueError:
                continue          # a truncated/garbled generation, not one we harvested
            key = norm(data.get("updated_question"))
            if key and key not in out:
                out[key] = {"system": rec["system"],
                            "user": rec["messages"][0]["content"],
                            "seed": rec["seed"],
                            "attempt": rec["attempt"],
                            "log_file": str(f),
                            # the original generation, for the side-by-side columns
                            "original": data}
    print(f"[logs] indexed {len(out)} generated questions from {len(files)} log files")
    return out


def patch_prompt(system: str) -> str:
    """Apply the required_reasoning_process change to a recovered system prompt."""
    assert FORMAT_INSERT not in system, "prompt already carries the new OUTPUT FORMAT key"
    assert system.count(FORMAT_ANCHOR) == 1, (
        f"expected 1 OUTPUT FORMAT anchor, found {system.count(FORMAT_ANCHOR)}"
    )
    assert system.count(RULES_ANCHOR) == 1, (
        f"expected 1 RULES anchor, found {system.count(RULES_ANCHOR)}"
    )
    system = system.replace(FORMAT_ANCHOR, FORMAT_ANCHOR + "\n" + FORMAT_INSERT)
    return system.replace(RULES_ANCHOR, RULES_ANCHOR + "\n" + RULES_INSERT)


# ---------------------------------------------------------------------------
# Generation
# ---------------------------------------------------------------------------

def generate_one(row: dict, call: dict, client, model: str, logger, pin_strategy: bool,
                 max_tokens: int, refusal_retries: int) -> dict:
    """One make-harder call under the patched prompt. No answering system is queried.

    Pinning a false-premise strategy while the new prompt also demands a spelled-out
    reasoning process can trip the model's refusal stop_reason. Retry, then fall back to
    an unpinned call so the row still yields a question; `strategy_pinned` records which.
    """
    system = patch_prompt(call["system"])
    base_user = call["user"]
    pin = STRATEGY_PIN.format(strategy=row["original_strategy"])

    attempts = ([(True, base_user + pin)] * (refusal_retries + 1) if pin_strategy else [])
    attempts.append((False, base_user))

    data, bucket, pinned, refusals = None, None, False, 0
    for pinned, user in attempts:
        try:
            raw, bucket = RP._call_llm(
                client, model=model, system=system,
                messages=[{"role": "user", "content": user}],
                max_tokens=max_tokens, logger=logger, seed=row["seed_question"],
                attempt=call["attempt"], purpose="harder",
            )
        # Narrow: only a call that produced no usable output falls back. RefusalError is
        # the pre-llm_client path; call_text now raises EmptyResponse for a refusal or an
        # otherwise empty completion, which is the same "nothing to parse" condition.
        except (RP.RefusalError, llm_client.EmptyResponse):
            refusals += 1
            continue
        data = RP.extract_json(raw)
        break
    assert data is not None, f"pos={row['position']}: refused even unpinned"

    return {
        **row,
        "strategy_pinned": pinned,
        "refusals": refusals,
        "new_question": data.get("updated_question", ""),
        "new_strategy": data.get("chosen_strategy", ""),
        "new_required_reasoning_process": RP.as_steps(
            data.get("required_reasoning_process")),
        "new_why_harder": data.get("why_harder", ""),
        "new_criterion": data.get("verification_criterion", ""),
        "new_brainstorming": data.get("brainstorming", ""),
        "cost_usd": bucket.cost_usd,
        "output_tokens": bucket.output_tokens,
        "log_file": call["log_file"],
    }


def attach_original_fields(results: list, calls: dict) -> list:
    """Add the pre-change generation's fields for side-by-side comparison.

    Derived from the logs at write time rather than stored in the generations cache, so
    rebuilding the workbook picks these up without re-running any calls.
    """
    for r in results:
        original = calls[norm(r["original_question"])]["original"]
        assert not original.get("required_reasoning_process"), (
            f"pos={r['position']}: original generation unexpectedly HAS a "
            f"required_reasoning_process; the n/a column would be wrong"
        )
        # verified.kept.json's criterion is the one annotators actually judged against;
        # for these rows it is byte-identical to the criterion as first generated.
        assert r["original_criterion"].strip() == (
            original.get("verification_criterion", "") or "").strip(), (
            f"pos={r['position']}: kept.json criterion differs from the logged one"
        )
        r["original_why_harder"] = original.get("why_harder", "")
        r["original_required_reasoning_process"] = NO_ORIGINAL_RRP
    return results


def generate_all(rows: list, calls: dict, args, cache_path: Path) -> list:
    """Generate every row, reusing any rows already checkpointed in cache_path."""
    done, legacy_stamped = {}, 0
    if cache_path.exists():
        done = {r["position"]: r for r in json.loads(cache_path.read_text())}
        # Rows cached before strategy_pinned/refusals existed came from the pinned-only
        # path, which had no fallback: reaching the cache at all means the pin held.
        legacy = [r for r in done.values() if "strategy_pinned" not in r]
        for r in legacy:
            r["strategy_pinned"], r["refusals"] = True, 0
        legacy_stamped = len(legacy)
        print(f"[gen] reusing {len(done)} cached generation(s) from {cache_path}"
              + (f" ({legacy_stamped} stamped as pinned)" if legacy else ""))

    def flush():
        ordered = [done[p] for p in sorted(done)]
        cache_path.write_text(json.dumps(ordered, indent=1, ensure_ascii=False))

    if legacy_stamped:
        flush()          # persist the migration so the cache matches the current schema

    todo = [r for r in rows if r["position"] not in done]
    if not todo:
        return [done[r["position"]] for r in rows]

    assert os.environ.get("ANTHROPIC_API_KEY"), "ANTHROPIC_API_KEY is not set"
    client = llm_client.make_client(args.model)
    logger = LockedRunLogger(cache_path.parent / "regenerate.jsonl", run_id="regenerate")
    lock = threading.Lock()

    print(f"[gen] {len(todo)} call(s) to make, concurrency {args.concurrency}")
    with ThreadPoolExecutor(max_workers=args.concurrency) as pool:
        futures = {
            pool.submit(generate_one, r, calls[norm(r["original_question"])], client,
                        args.model, logger, not args.no_pin_strategy, args.max_tokens,
                        args.refusal_retries): r
            for r in todo
        }
        for fut in as_completed(futures):
            row = futures[fut]
            result = fut.result()          # let a failure surface; finished rows are cached
            with lock:
                done[row["position"]] = result
                flush()
            print(f"  [{len(done)}/{len(rows)}] pos={row['position']:2d} "
                  f"{row['annotator']:9s} {result['new_question'][:70]}")

    return [done[r["position"]] for r in rows]


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

# Appended to the right of the original columns; the annotation columns keep their
# original positions so the sheet is filled in exactly as before.
REFERENCE_COLUMNS = [
    "Verification Criterion (regenerated)",
    "Required Reasoning Process (regenerated)",
    "Why Harder (regenerated)",
    "Strategy (pinned, from original run)",
    "Original Seed Question",
    "Previously Annotated Question (for comparison)",
    "Verification Criterion (original)",
    "Why Harder (original)",
    "Required Reasoning Process (original — n/a, field postdates this run)",
    "Strategy Control",
]

# The original run predates required_reasoning_process, so there is nothing to show. A
# literal marker beats an empty cell, which reads as an oversight.
NO_ORIGINAL_RRP = "n/a"


def write_workbook(results: list, xlsx: Path, sheet: str, out: Path,
                   annotators: list) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font

    src = openpyxl.load_workbook(xlsx, read_only=True)[sheet]
    header = [h for h in next(src.iter_rows(values_only=True))]
    while header and header[-1] is None:
        header.pop()
    a_col = col_index(header, "Annotator")
    q_col = col_index(header, "Final Question")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header + REFERENCE_COLUMNS)

    # Same grouping as the source sheet: annotator blocks, positions ascending within.
    ordered = sorted(results, key=lambda r: (annotators.index(r["annotator"]),
                                             r["position"]))
    for r in ordered:
        row = [None] * len(header)
        row[a_col] = r["annotator"]
        row[q_col] = r["new_question"]
        ws.append(row + [
            r["new_criterion"],
            # one numbered step per line, so the cell stays readable when wrapped
            "\n".join(f"{i}. {s}" for i, s in
                      enumerate(RP.as_steps(r["new_required_reasoning_process"]), 1)),
            r["new_why_harder"],
            r["original_strategy"],
            r["seed_question"],
            r["original_question"],
            r["original_criterion"],
            r["original_why_harder"],
            r["original_required_reasoning_process"],
            "pinned" if r["strategy_pinned"]
            else f"UNPINNED after {r['refusals']} refusal(s) — model re-chose",
        ])

    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    widths = {a_col: 12, q_col: 70}
    for i in range(len(header) + len(REFERENCE_COLUMNS)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = (
            widths.get(i, 45)
        )
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"[out] wrote {len(ordered)} rows to {out}")


def main() -> int:
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    p.add_argument("--xlsx", required=True,
                   help="Source annotation workbook (read only; never written).")
    p.add_argument("--sheet", default="random assignment",
                   help="Sheet in --xlsx whose row set and columns to mirror.")
    p.add_argument("--kept", required=True,
                   help="verified.kept.json the sheet was drawn from.")
    p.add_argument("--logs", required=True,
                   help="Run root searched recursively for run-sample_*.jsonl logs.")
    p.add_argument("--out", required=True, help="Output .xlsx path.")
    p.add_argument("--generations-json", required=True,
                   help="Checkpoint file for the raw generations; reused on re-runs.")
    p.add_argument("--annotators", default=",".join(DEFAULT_ANNOTATORS),
                   help="Comma-separated deal order.")
    p.add_argument("--n-rows", type=int, default=30, help="Rows to reproduce.")
    p.add_argument("--model", default=RP.CLAUDE_MODEL)
    p.add_argument("--max-tokens", type=int, default=RP.MAKE_HARDER_MAX_TOKENS,
                   help=f"Matches research_pipeline's make-harder call "
                        f"(default: {RP.MAKE_HARDER_MAX_TOKENS}).")
    p.add_argument("--concurrency", type=int, default=6)
    p.add_argument("--no-pin-strategy", action="store_true",
                   help="Let the model choose its own strategy instead of pinning the "
                        "one the original question used.")
    p.add_argument("--refusal-retries", type=int, default=2,
                   help="Retries of a pinned call that ends in stop_reason='refusal' "
                        "before falling back to an unpinned call (default: 2).")
    args = p.parse_args()

    xlsx = Path(args.xlsx).expanduser()
    out = Path(args.out).expanduser()
    assert xlsx.resolve() != out.resolve(), "--out would overwrite the source workbook"

    annotators = [a.strip() for a in args.annotators.split(",") if a.strip()]
    rows = build_assignment(Path(args.kept).expanduser(), annotators, args.n_rows)
    verify_against_sheet(rows, xlsx, args.sheet)

    calls = index_harder_calls(Path(args.logs).expanduser())
    missing = [r["original_question"] for r in rows
               if norm(r["original_question"]) not in calls]
    assert not missing, f"no originating generation call found for: {missing}"

    results = attach_original_fields(
        generate_all(rows, calls, args, Path(args.generations_json).expanduser()), calls
    )
    total = sum(r["cost_usd"] for r in results)
    print(f"[gen] {len(results)} question(s); cost ${total:.4f}")

    write_workbook(results, xlsx, args.sheet, out, annotators)
    return 0


if __name__ == "__main__":
    sys.exit(main())
